__all__ = ["XLSXVisitor"]
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from typing import Any

from tg_bot.models import User, TestResult, TestResponse, KnowledgeAssesment, Excursion, Administrator
from tg_bot.utils.interfaces import IModelVisitor


class XLSXVisitor(IModelVisitor):
    """A visitor class that exports model data to XLSX format.

    This class implements the IModelVisitor interface to visit various models
    and export their data to Excel spreadsheets. Each model type gets its own
    worksheet with appropriate column headers and data.

    :ivar _workbook: The Excel workbook being created
    """

    def __init__(self):
        """Initializes a new XLSXVisitor instance.

        Creates a new workbook and removes the default sheet.
        :return: None
        """
        self._workbook: Workbook = Workbook()
        self._workbook.remove(self._workbook.active)

    async def visit_administrator(self) -> None:
        """Visits administrator models and exports their data to an Excel sheet.

        Creates a sheet named "Администратор" with ID and permission data.
        :return: None
        """
        sheet: Worksheet = self._workbook.create_sheet(title="Администратор")
        annotations = ["ID администратора", "Может управлять администраторами"]
        objects = await Administrator.all().values_list("id", "can_manipulate_administrator_table")
        await self.export_data_to_xlsx_sheet(annotations, objects, sheet)

    async def visit_bot_user_profile(self) -> None:
        """Placeholder for bot user profile export functionality.

        :return: None
        """
        pass

    async def visit_excursion(self) -> None:
        """Visits excursion models and exports their data to Excel sheets.

        Creates two sheets:
        1. "Экскурсия" with address and date information
        2. "Купленные экскурсии" with user IDs and excursion dates
        :return: None
        """
        sheet: Worksheet = self._workbook.create_sheet(title="Экскурсия")
        annotations = ["Адрес", "Дата"]
        objects = await Excursion.all()
        objects = tuple([(obj.address, obj.date) for obj in objects])
        await self.export_data_to_xlsx_sheet(annotations, objects, sheet)
        sheet = self._workbook.create_sheet(title="Купленные экскурсии")
        annotations = ["ID пользователя", "Дата экскурсии"]
        excursions = await Excursion.all()
        objects = []
        for excursion in excursions:
            users = [await profile.get_user() for profile in await excursion.bot_user_profiles.all()]
            for user in users:
                objects.append((user.id, excursion.date))
        objects = tuple(objects)
        await self.export_data_to_xlsx_sheet(annotations, objects, sheet)

    async def visit_knowledge_assesment(self) -> None:
        """
        Visits knowledge assessment models and exports their data to Excel sheets.
        Creates two sheets:
        1. "Проверка знаний" with webinar link and date information
        2. "Купленные проверки знаний" with user IDs and assessment dates
        """
        sheet: Worksheet = self._workbook.create_sheet(title="Проверка знаний")
        annotations = ["Ссылка", "Дата"]
        objects = await KnowledgeAssesment.all()
        objects = tuple([(obj.webinar_link, obj.date) for obj in objects])
        await self.export_data_to_xlsx_sheet(annotations, objects, sheet)
        sheet = self._workbook.create_sheet(title="Купленные проверки знаний")
        annotations = ["ID пользователя", "Дата проверки знаний"]
        assesments = await KnowledgeAssesment.all()
        objects = []
        for assesment in assesments:
            users = [await profile.get_user() for profile in await assesment.bot_user_profiles.all()]
            for user in users:
                objects.append((user.id, assesment.date))
        objects = tuple(objects)
        await self.export_data_to_xlsx_sheet(annotations, objects, sheet)

    async def visit_order(self) -> None:
        """
        Placeholder for order export functionality.
        """
        pass

    async def visit_school_grade(self) -> None:
        """
        Placeholder for school grade export functionality.
        """
        pass

    async def visit_test_answer(self) -> None:
        """
        Placeholder for test answer export functionality.
        """
        pass

    async def visit_test_question(self) -> None:
        """
        Placeholder for test question export functionality.
        """
        pass

    async def visit_test_response(self) -> None:
        """
        Placeholder for test response export functionality.
        """
        pass

    async def visit_test_result(self) -> None:
        """
        Visits test result models and exports their data to an Excel sheet.
        Creates a sheet named "Результаты тестирования" with user IDs, grades,
        and all questions and answers from the test.
        """
        sheet: Worksheet = self._workbook.create_sheet(title="Результаты тестирования")
        annotations = ["ID пользователя", "Класс"]
        test_results = await TestResult.all()
        users = [await (await obj.get_bot_user_profile()).get_user() for obj in test_results]
        objects: list[Any | None] = [None] * len(test_results)
        max_count_of_questions = -1
        for i in range(len(test_results)):
            objects[i] = [users[i].id, (await test_results[i].grade).grade]
            responses: list[TestResponse] = await test_results[i].test_responses.all()
            count_of_questions = 0
            for response in responses:
                count_of_questions += 1
                objects[i].append((await response.question).description)
                objects[i].append((await response.answer).description)
            max_count_of_questions = max(max_count_of_questions, count_of_questions)
        annotations += ["Вопрос", "Ответ"] * max_count_of_questions
        await self.export_data_to_xlsx_sheet(annotations, tuple(objects), sheet)

    async def visit_user(self) -> None:
        """
        Visits user models and exports their data to an Excel sheet.
        Creates a sheet named "Пользователи" with ID, name, and phone information.
        """
        sheet = self._workbook.create_sheet(title="Пользователи")
        annotations = ["ID Пользователя", "ФИО", "Телефон"]
        objects = await User.all().values_list("id", "name_last_name_surname", "phone_number")
        await self.export_data_to_xlsx_sheet(annotations, objects, sheet)

    @property
    def workbook(self) -> Workbook:
        """
        Property to access the generated workbook.
        
        Returns:
            Workbook: The Excel workbook containing all exported data
        """
        return self._workbook

    @staticmethod
    async def export_data_to_xlsx_sheet(annotations: list[str], objects: tuple[tuple[Any, ...], ...], sheet: Worksheet) -> None:
        """
        Exports data to an Excel worksheet.
        
        Args:
            annotations (list[str]): Column headers for the worksheet
            objects (tuple[tuple[Any, ...], ...]): Data to export, where each inner tuple represents a row
            sheet (Worksheet): The worksheet to export the data to
        """
        for i, annotation in enumerate(annotations):
            sheet[f"{chr(ord('A') + i)}1"] = annotation
        for i, obj in enumerate(objects, start=2):
            for j, field in enumerate(obj):
                sheet[f"{chr(ord('A') + j)}{i}"] = field